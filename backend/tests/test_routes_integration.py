import asyncio
from collections.abc import AsyncGenerator
from io import BytesIO
from urllib.parse import parse_qs, urlparse

import httpx
from fastapi import FastAPI
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from app.api.router import router as api_router
from app.api.v1 import admin as admin_api
from app.api.v1 import student as student_api
from app.api.v1 import teacher as teacher_api
from app.api.v1 import evidence as evidence_api
from app.core.security import create_access_token, hash_password
from app.db.base import Base
from app.db.models import Chapter, Course, StudentMilestone, Subject, TeacherProfile, User
from app.db.models import MilestoneStatus, UserRole
from app.db.session import get_db
from app.services.ai import advisor


def test_routes_cover_student_teacher_admin_and_public_flows(
    monkeypatch, tmp_path,
) -> None:
    monkeypatch.setattr(advisor, "GROQ_API_KEY", "")
    monkeypatch.setattr(advisor, "GEMINI_API_KEY", "")
    monkeypatch.setattr(evidence_api, "UPLOAD_DIR", tmp_path)
    review_notifications: list[tuple[str, str, str, str, str | None]] = []
    student_invitations: list[tuple[str, str, str]] = []

    async def record_student_invitation(
        email: str, name: str, invitation_url: str
    ) -> None:
        student_invitations.append((email, name, invitation_url))

    monkeypatch.setattr(admin_api, "ensure_email_configured", lambda: None)
    monkeypatch.setattr(
        admin_api, "dispatch_student_invitation", record_student_invitation
    )

    async def record_review_notification(
        email: str,
        name: str,
        title: str,
        decision: str,
        feedback: str | None,
    ) -> None:
        review_notifications.append((email, name, title, decision, feedback))

    monkeypatch.setattr(
        teacher_api, "dispatch_review_notification", record_review_notification
    )

    async def github_profile(username: str) -> dict:
        assert username == "sample-student"
        return {
            "login": username,
            "name": "Sample Student",
            "bio": None,
            "public_repos": 3,
            "followers": 2,
            "following": 1,
            "html_url": "https://github.com/sample-student",
        }

    monkeypatch.setattr(student_api, "get_github_profile", github_profile)

    async def exercise_routes() -> None:
        engine = create_async_engine(
            "sqlite+aiosqlite:///:memory:",
            poolclass=StaticPool,
        )
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)

        async with session_factory() as session:
            course = Course(name="Computer Science", duration_years=4)
            session.add(course)
            await session.flush()
            subject = Subject(
                course_id=course.id,
                code="CS101",
                name="Programming",
                semester=1,
            )
            session.add(subject)
            await session.flush()
            session.add_all(
                [
                    Chapter(
                        subject_id=subject.id,
                        sequence_order=1,
                        title="Foundations",
                        description="Learn core programming.",
                    ),
                    Chapter(
                        subject_id=subject.id,
                        sequence_order=2,
                        title="Data structures",
                        description="Practice common data structures.",
                    ),
                ]
            )
            admin = User(
                email="admin@example.com",
                hashed_password=hash_password("valid-admin-password"),
                role=UserRole.ADMIN,
            )
            teacher = User(
                email="teacher@example.com",
                hashed_password=hash_password("valid-teacher-password"),
                role=UserRole.TEACHER,
            )
            teacher.teacher_profile = TeacherProfile(
                first_name="Terry",
                last_name="Teacher",
                department="Computer Science",
            )
            teacher.assigned_courses = [course]
            session.add_all([admin, teacher])
            await session.commit()
            course_id = course.id
            admin_id = admin.id
            teacher_id = teacher.id

        application = FastAPI()
        application.include_router(api_router)

        async def override_get_db() -> AsyncGenerator[AsyncSession, None]:
            async with session_factory() as session:
                try:
                    yield session
                    await session.commit()
                except Exception:
                    await session.rollback()
                    raise

        application.dependency_overrides[get_db] = override_get_db
        tokens = {
            "admin": create_access_token(str(admin_id)),
            "teacher": create_access_token(str(teacher_id)),
        }

        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=application),
            base_url="http://test",
        ) as client:
            response = await client.get("/api/v1/courses")
            assert response.status_code == 200
            assert response.json()[0]["id"] == course_id

            response = await client.get("/api/v1/announcements")
            assert response.status_code == 200
            assert response.json()[0]["id"] == "welcome"

            response = await client.get("/api/v1/student/dashboard")
            assert response.status_code == 401

            response = await client.post(
                "/api/v1/auth/register",
                json={
                    "email": " Student@Example.com ",
                    "password": "student-password-123",
                    "first_name": "Sam",
                    "last_name": "Student",
                    "course_id": course_id,
                },
            )
            assert response.status_code == 201
            student_token = response.json()["access_token"]
            student_id = response.json()["user"]["id"]
            assert response.json()["user"]["email"] == "student@example.com"
            student_headers = {"Authorization": f"Bearer {student_token}"}

            response = await client.post(
                "/api/v1/auth/register",
                json={
                    "email": "student@example.com",
                    "password": "student-password-123",
                    "first_name": "Sam",
                    "last_name": "Student",
                    "course_id": course_id,
                },
            )
            assert response.status_code == 409

            response = await client.post(
                "/api/v1/auth/login",
                json={
                    "email": "STUDENT@example.com",
                    "password": "student-password-123",
                },
            )
            assert response.status_code == 200
            assert response.json()["user"]["email"] == "student@example.com"

            response = await client.get(
                "/api/v1/auth/me",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.json()["first_name"] == "Sam"

            response = await client.get(
                "/api/v1/student/dashboard",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.json()["stats"]["total_milestones"] == 2
            response = await client.get(
                "/api/v1/student/telemetry/github/sample-student",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.json()["public_repos"] == 3

            monkeypatch.setattr(
                student_api, "read_authorization_state", lambda state: student_id
            )

            async def linked_in_profile(code: str) -> dict[str, str | None]:
                assert code == "test-code"
                return {
                    "sub": "linkedin-id",
                    "name": "Sample Student",
                    "picture": "https://media.example.com/profile.jpg",
                    "given_name": "Sample",
                    "family_name": "Student",
                }

            monkeypatch.setattr(
                student_api, "exchange_code_for_profile", linked_in_profile
            )
            response = await client.get(
                "/api/v1/student/telemetry/linkedin/callback",
                params={"code": "test-code", "state": "signed-state"},
                follow_redirects=False,
            )
            assert response.status_code == 303
            assert "linkedin=connected" in response.headers["location"]
            response = await client.get(
                "/api/v1/student/telemetry/linkedin/profile",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.json() == {
                "name": "Sample Student",
                "picture_url": "https://media.example.com/profile.jpg",
                "connected": True,
            }

            response = await client.get(
                "/api/v1/student/milestones",
                headers=student_headers,
            )
            assert response.status_code == 200
            milestones = response.json()
            assert [milestone["status"] for milestone in milestones] == [
                MilestoneStatus.IN_PROGRESS.value,
                MilestoneStatus.LOCKED.value,
            ]
            first_milestone_id = milestones[0]["id"]
            second_milestone_id = milestones[1]["id"]

            response = await client.post(
                "/api/v1/student/evidence",
                headers=student_headers,
                files={"file": ("proof.pdf", b"%PDF-1.4\nstudent proof", "application/pdf")},
            )
            assert response.status_code == 201
            milestone_evidence_id = response.json()["id"]
            response = await client.get(
                f"/api/v1/student/evidence/{milestone_evidence_id}",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.content.startswith(b"%PDF-")

            response = await client.post(
                "/api/v1/student/evidence",
                headers=student_headers,
                files={"file": ("malware.txt", b"not a PDF", "application/pdf")},
            )
            assert response.status_code == 415
            assert len(list(tmp_path.iterdir())) == 1

            response = await client.post(
                f"/api/v1/student/milestones/{second_milestone_id}/submit",
                headers=student_headers,
                json={"summary": "I finished the locked chapter."},
            )
            assert response.status_code == 409

            response = await client.post(
                f"/api/v1/student/milestones/{first_milestone_id}/submit",
                headers=student_headers,
                json={
                    "summary": "I completed the first programming exercise.",
                    "evidence_file_id": milestone_evidence_id,
                },
            )
            assert response.status_code == 200
            assert response.json()["status"] == MilestoneStatus.PENDING_VERIFICATION.value

            response = await client.post(
                f"/api/v1/student/milestones/{first_milestone_id}/submit",
                headers=student_headers,
                json={"summary": "I submitted this a second time."},
            )
            assert response.status_code == 409

            response = await client.get(
                "/api/v1/teacher/dashboard",
                headers=student_headers,
            )
            assert response.status_code == 403

            teacher_headers = {"Authorization": f"Bearer {tokens['teacher']}"}
            response = await client.get(
                "/api/v1/teacher/dashboard",
                headers=teacher_headers,
            )
            assert response.status_code == 200
            assert response.json()["stats"]["pending_reviews"] == 1

            response = await client.get(
                "/api/v1/teacher/submissions",
                headers=teacher_headers,
            )
            assert response.status_code == 200
            assert response.json()[0]["id"] == first_milestone_id
            assert response.json()[0]["evidence_file_id"] == milestone_evidence_id
            response = await client.get(
                f"/api/v1/student/evidence/{milestone_evidence_id}",
                headers=teacher_headers,
            )
            assert response.status_code == 200

            response = await client.post(
                f"/api/v1/teacher/submissions/{first_milestone_id}/review",
                headers=teacher_headers,
                json={"decision": "needs_revision", "feedback": "Add more detail."},
            )
            assert response.status_code == 200
            assert response.json()["status"] == MilestoneStatus.IN_PROGRESS.value

            response = await client.post(
                f"/api/v1/student/milestones/{first_milestone_id}/submit",
                headers=student_headers,
                json={"summary": "I added details to my programming exercise."},
            )
            assert response.status_code == 200

            response = await client.post(
                f"/api/v1/teacher/submissions/{first_milestone_id}/review",
                headers=teacher_headers,
                json={"decision": "verified", "feedback": ""},
            )
            assert response.status_code == 200
            assert response.json()["status"] == MilestoneStatus.VERIFIED.value

            response = await client.get(
                "/api/v1/student/milestones",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert [milestone["status"] for milestone in response.json()] == [
                MilestoneStatus.VERIFIED.value,
                MilestoneStatus.IN_PROGRESS.value,
            ]

            response = await client.post(
                f"/api/v1/teacher/submissions/{first_milestone_id}/review",
                headers=teacher_headers,
                json={"decision": "unknown", "feedback": ""},
            )
            assert response.status_code == 422

            admin_headers = {"Authorization": f"Bearer {tokens['admin']}"}
            response = await client.get(
                "/api/v1/teacher/academic-options",
                headers=student_headers,
            )
            assert response.status_code == 403

            response = await client.get(
                "/api/v1/teacher/academic-options",
                headers=teacher_headers,
            )
            assert response.status_code == 200
            academic_options = response.json()
            academic_student_id = academic_options["students"][0]["id"]
            academic_subject_id = academic_options["subjects"][0]["id"]

            attendance_entry = {
                "student_id": academic_student_id,
                "subject_id": academic_subject_id,
                "attendance_date": "2026-09-01",
                "status": "present",
            }
            response = await client.post(
                "/api/v1/teacher/attendance",
                headers=teacher_headers,
                json=attendance_entry,
            )
            assert response.status_code == 200
            attendance_entry["status"] = "absent"
            response = await client.post(
                "/api/v1/teacher/attendance",
                headers=teacher_headers,
                json=attendance_entry,
            )
            assert response.status_code == 200
            response = await client.post(
                "/api/v1/teacher/attendance",
                headers=teacher_headers,
                json={
                    **attendance_entry,
                    "attendance_date": "2026-09-02",
                    "status": "present",
                },
            )
            assert response.status_code == 200
            response = await client.post(
                "/api/v1/teacher/marks",
                headers=teacher_headers,
                json={
                    "student_id": academic_student_id,
                    "subject_id": academic_subject_id,
                    "title": "Quiz 1",
                    "score": 8,
                    "maximum_score": 10,
                    "assessment_date": "2026-09-02",
                },
            )
            assert response.status_code == 200
            response = await client.post(
                "/api/v1/teacher/marks",
                headers=teacher_headers,
                json={
                    "student_id": academic_student_id,
                    "subject_id": academic_subject_id,
                    "title": "Invalid quiz",
                    "score": 11,
                    "maximum_score": 10,
                    "assessment_date": "2026-09-02",
                },
            )
            assert response.status_code == 422
            response = await client.get(
                "/api/v1/student/academic-records",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.json()["attendance"][0]["percentage"] == 50.0
            assert response.json()["marks"][0]["percentage"] == 80.0
            response = await client.get(
                "/api/v1/admin/reports/academics",
                headers=admin_headers,
            )
            assert response.status_code == 200
            assert response.json()["attendance"][0]["attendance_percentage"] == 50.0
            assert response.json()["assessments"][0]["average_percentage"] == 80.0

            response = await client.post(
                "/api/v1/student/activities",
                headers=student_headers,
                json={
                    "category": "internship",
                    "title": "Software engineering internship",
                    "description": "Built and tested a small service during an internship.",
                    "issuer": "Example Labs",
                    "achieved_on": "2026-06-15",
                },
            )
            assert response.status_code == 422

            response = await client.post(
                "/api/v1/student/evidence",
                headers=student_headers,
                files={"file": ("internship.exe", b"%PDF-1.7\ninternship proof", "application/pdf")},
            )
            assert response.status_code == 201
            activity_evidence_id = response.json()["id"]
            assert response.json()["filename"] == "internship.pdf"
            stored_file = next(tmp_path.iterdir())
            assert stored_file.stat().st_mode & 0o777 == 0o600
            response = await client.post(
                "/api/v1/student/activities",
                headers=student_headers,
                json={
                    "category": "internship",
                    "title": "Software engineering internship",
                    "description": "Built and tested a small service during an internship.",
                    "issuer": "Example Labs",
                    "achieved_on": "2026-06-15",
                    "evidence_file_id": activity_evidence_id,
                },
            )
            assert response.status_code == 201
            activity = response.json()
            assert activity["status"] == "pending_verification"
            response = await client.get(
                f"/api/v1/student/evidence/{activity_evidence_id}",
                headers=teacher_headers,
            )
            assert response.status_code == 200
            assert response.headers["content-disposition"].endswith(
                'filename="internship.pdf"'
            )

            response = await client.get(
                "/api/v1/teacher/activities",
                headers=student_headers,
            )
            assert response.status_code == 403
            response = await client.get(
                "/api/v1/student/portfolio",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.json()["activities"] == []

            response = await client.get(
                "/api/v1/teacher/activities",
                headers=teacher_headers,
            )
            assert response.status_code == 200
            assert response.json()[0]["id"] == activity["id"]
            response = await client.post(
                f"/api/v1/teacher/activities/{activity['id']}/review",
                headers=teacher_headers,
                json={"decision": "verified", "feedback": "Evidence checked."},
            )
            assert response.status_code == 200
            assert response.json()["status"] == "verified"
            assert review_notifications[-1] == (
                "student@example.com",
                "Sam Student",
                "Software engineering internship",
                "verified",
                "Evidence checked.",
            )
            response = await client.put(
                f"/api/v1/student/activities/{activity['id']}",
                headers=student_headers,
                json={
                    "category": "internship",
                    "title": "Changed verified title",
                    "description": "Attempting to edit a verified achievement.",
                    "proof_url": "https://example.com/proof",
                },
            )
            assert response.status_code == 409

            response = await client.get(
                "/api/v1/student/portfolio",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.json()["activities"][0]["title"] == (
                "Software engineering internship"
            )
            assert response.json()["activities"][0]["evidence_file_id"] == activity_evidence_id
            assert "email" not in response.json()["student"]
            assert response.json()["academic_milestones"][0]["title"] == "Foundations"
            assert response.json()["share_enabled"] is False
            response = await client.get(
                "/api/v1/student/portfolio/resume",
                headers=student_headers,
            )
            assert response.status_code == 200
            assert response.headers["content-disposition"].endswith(
                'filename="verisphere-resume.html"'
            )
            assert "Foundations" in response.text
            assert "Software engineering internship" in response.text
            response = await client.post(
                "/api/v1/student/portfolio/share",
                headers=student_headers,
            )
            assert response.status_code == 200
            share_path = response.json()["share_url"]
            share_token = share_path.rsplit("/", 1)[-1]
            response = await client.get(f"/api/v1/portfolios/{share_token}")
            assert response.status_code == 200
            assert response.json()["activities"][0]["title"] == (
                "Software engineering internship"
            )
            response = await client.get(
                f"/api/v1/portfolios/{share_token}/resume"
            )
            assert response.status_code == 200
            assert "Software engineering internship" in response.text
            response = await client.get(
                f"/api/v1/student/evidence/shared/{share_token}/{activity_evidence_id}",
            )
            assert response.status_code == 200

            response = await client.get(
                "/api/v1/admin/reports/activities",
                headers=admin_headers,
            )
            assert response.status_code == 200
            assert response.json()["total_activities"] == 1
            assert response.json()["verified_activities"] == 1
            assert response.json()["by_category"] == [
                {"category": "internship", "count": 1}
            ]
            response = await client.get(
                "/api/v1/admin/reports/activities/export.xlsx",
                headers=admin_headers,
            )
            assert response.status_code == 200
            assert response.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            activity_workbook = load_workbook(BytesIO(response.content), read_only=True)
            assert activity_workbook.sheetnames == ["Summary", "Categories", "Statuses"]
            assert activity_workbook["Categories"]["A2"].value == "internship"
            response = await client.get(
                "/api/v1/admin/reports/academics/export.xlsx",
                headers=admin_headers,
            )
            assert response.status_code == 200
            academic_workbook = load_workbook(BytesIO(response.content), read_only=True)
            assert academic_workbook.sheetnames == [
                "Summary",
                "Attendance",
                "Assessments",
            ]
            response = await client.get(
                "/api/v1/admin/reports/activities",
                headers=teacher_headers,
            )
            assert response.status_code == 403

            response = await client.delete(
                "/api/v1/student/portfolio/share",
                headers=student_headers,
            )
            assert response.status_code == 204
            response = await client.get(f"/api/v1/portfolios/{share_token}")
            assert response.status_code == 404
            response = await client.get(
                f"/api/v1/student/evidence/shared/{share_token}/{activity_evidence_id}",
            )
            assert response.status_code == 404

            response = await client.get(
                "/api/v1/admin/overview",
                headers=admin_headers,
            )
            assert response.status_code == 200
            assert response.json()["stats"]["students"] == 1
            assert response.json()["stats"]["teachers"] == 1

            response = await client.get(
                "/api/v1/admin/users",
                headers=admin_headers,
            )
            assert response.status_code == 200
            assert len(response.json()) == 3

            response = await client.post(
                "/api/v1/admin/teachers",
                headers=admin_headers,
                json={
                    "email": "new.teacher@example.com",
                    "password": "new-teacher-password",
                    "first_name": "New",
                    "last_name": "Teacher",
                    "department": "Computer Science",
                    "course_ids": [course_id],
                },
            )
            assert response.status_code == 201
            new_teacher_id = response.json()["id"]
            new_teacher_headers = {
                "Authorization": f"Bearer {create_access_token(str(new_teacher_id))}"
            }
            response = await client.get(
                "/api/v1/admin/users",
                headers=admin_headers,
            )
            assert response.status_code == 200
            assigned_teacher = next(
                account for account in response.json()
                if account["id"] == new_teacher_id
            )
            assert assigned_teacher["courses"] == [
                {"id": course_id, "name": "Computer Science"}
            ]
            response = await client.get(
                "/api/v1/teacher/academic-options",
                headers=new_teacher_headers,
            )
            assert response.status_code == 200
            assert len(response.json()["students"]) == 1
            assert [subject["code"] for subject in response.json()["subjects"]] == [
                "CS101"
            ]
            response = await client.put(
                f"/api/v1/admin/teachers/{new_teacher_id}/courses",
                headers=admin_headers,
                json={"course_ids": []},
            )
            assert response.status_code == 200
            response = await client.get(
                "/api/v1/teacher/dashboard",
                headers=new_teacher_headers,
            )
            assert response.status_code == 200
            assert response.json()["stats"] == {
                "pending_reviews": 0,
                "pending_activity_reviews": 0,
                "students": 0,
            }
            response = await client.get(
                "/api/v1/teacher/academic-options",
                headers=new_teacher_headers,
            )
            assert response.status_code == 200
            assert response.json() == {"students": [], "subjects": []}
            response = await client.post(
                "/api/v1/teacher/attendance",
                headers=new_teacher_headers,
                json={
                    "student_id": 1,
                    "subject_id": 1,
                    "attendance_date": "2026-10-17",
                    "status": "present",
                },
            )
            assert response.status_code == 404
            response = await client.post(
                "/api/v1/teacher/submissions/1/review",
                headers=new_teacher_headers,
                json={"decision": "verified"},
            )
            assert response.status_code == 404
            response = await client.post(
                "/api/v1/teacher/activities/1/review",
                headers=new_teacher_headers,
                json={"decision": "verified"},
            )
            assert response.status_code == 404

            response = await client.post(
                "/api/v1/admin/courses",
                headers=admin_headers,
                json={"name": "Advanced Course", "duration_years": 3},
            )
            assert response.status_code == 201
            response = await client.post(
                "/api/v1/admin/courses",
                headers=admin_headers,
                json={"name": "  Advanced Course  ", "duration_years": 3},
            )
            assert response.status_code == 409

            response = await client.post(
                "/api/v1/assistant/chat",
                headers=student_headers,
                json={"messages": [{"role": "user", "content": "What next?"}]},
            )
            assert response.status_code == 503

            response = await client.post(
                "/api/v1/assistant/chat",
                headers=admin_headers,
                json={"messages": [{"role": "user", "content": "What next?"}]},
            )
            assert response.status_code == 403
            response = await client.post(
                "/api/v1/assistant/feedback",
                headers=student_headers,
                json={"provider": "Groq", "helpful": True},
            )
            assert response.status_code == 201
            response = await client.get(
                "/api/v1/admin/assistant/feedback",
                headers=admin_headers,
            )
            assert response.status_code == 200
            assert response.json()["providers"] == [
                {
                    "provider": "Groq",
                    "responses_rated": 1,
                    "helpful": 1,
                    "helpfulness_percentage": 100.0,
                }
            ]

            response = await client.get(
                "/api/v1/admin/students/import-template.xlsx",
                headers=admin_headers,
            )
            assert response.status_code == 200
            template = load_workbook(BytesIO(response.content), read_only=True)
            assert tuple(cell.value for cell in template.active[1]) == (
                "email",
                "first_name",
                "last_name",
                "course_id",
            )
            template.close()

            workbook = Workbook()
            workbook.active.append(
                ["email", "first_name", "last_name", "course_id"]
            )
            workbook.active.append(
                ["new-student@example.com", "New", "Student", course_id]
            )
            workbook_bytes = BytesIO()
            workbook.save(workbook_bytes)
            workbook.close()
            response = await client.post(
                "/api/v1/admin/students/import",
                headers=admin_headers,
                files={
                    "file": (
                        "students.xlsx",
                        workbook_bytes.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert response.status_code == 202
            assert response.json() == {"invited": 1}
            assert student_invitations[0][:2] == (
                "new-student@example.com",
                "New Student",
            )
            invitation_token = parse_qs(
                urlparse(student_invitations[0][2]).query
            )["token"][0]
            response = await client.post(
                "/api/v1/auth/register",
                json={
                    "email": "new-student@example.com",
                    "password": "another-password-123",
                    "first_name": "New",
                    "last_name": "Student",
                    "course_id": course_id,
                },
            )
            assert response.status_code == 409
            response = await client.post(
                "/api/v1/auth/invitations/accept",
                json={
                    "token": invitation_token,
                    "password": "new-student-password",
                },
            )
            assert response.status_code == 201
            response = await client.post(
                "/api/v1/auth/invitations/accept",
                json={
                    "token": invitation_token,
                    "password": "another-password-123",
                },
            )
            assert response.status_code == 409

        await engine.dispose()

    asyncio.run(exercise_routes())
