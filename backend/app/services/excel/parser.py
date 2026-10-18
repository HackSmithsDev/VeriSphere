from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException


MAX_STUDENT_IMPORT_ROWS = 50
STUDENT_IMPORT_COLUMNS = ("email", "first_name", "last_name", "course_id")


class SpreadsheetValidationError(ValueError):
    def __init__(self, errors: list[str]) -> None:
        self.errors = errors
        super().__init__("Student import workbook contains invalid rows.")


@dataclass(frozen=True)
class StudentImportRow:
    email: str
    first_name: str
    last_name: str
    course_id: int
    row_number: int


def parse_student_workbook(content: bytes) -> list[StudentImportRow]:
    if not content:
        raise SpreadsheetValidationError(["The uploaded workbook is empty."])
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise SpreadsheetValidationError(
            ["Upload a valid .xlsx workbook."]
        ) from error

    try:
        if len(workbook.worksheets) != 1:
            raise SpreadsheetValidationError(
                ["The workbook must contain exactly one worksheet."]
            )
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise SpreadsheetValidationError(["The workbook has no header row."])
        headers = [
            str(value).strip().lower() if value is not None else ""
            for value in header_row
        ]
        if len(headers) != len(STUDENT_IMPORT_COLUMNS) or tuple(headers) != STUDENT_IMPORT_COLUMNS:
            raise SpreadsheetValidationError(
                [
                    "Use exactly these columns in order: "
                    + ", ".join(STUDENT_IMPORT_COLUMNS)
                    + "."
                ]
            )

        imported: list[StudentImportRow] = []
        errors: list[str] = []
        seen_emails: set[str] = set()
        for row_number, values in enumerate(rows, start=2):
            if all(value is None for value in values):
                continue
            if len(imported) + len(errors) >= MAX_STUDENT_IMPORT_ROWS:
                errors.append(
                    f"Import is limited to {MAX_STUDENT_IMPORT_ROWS} non-empty rows."
                )
                break
            if len(values) != len(STUDENT_IMPORT_COLUMNS):
                errors.append(f"Row {row_number}: expected four columns.")
                continue
            if any(isinstance(value, str) and value.startswith("=") for value in values):
                errors.append(f"Row {row_number}: formulas are not allowed.")
                continue
            email_value, first_value, last_value, course_value = values
            email = str(email_value or "").strip().lower()
            first_name = str(first_value or "").strip()
            last_name = str(last_value or "").strip()
            try:
                email = validate_email(email, check_deliverability=False).normalized
            except EmailNotValidError:
                errors.append(f"Row {row_number}: enter a valid email address.")
                continue
            if not 1 <= len(first_name) <= 100 or not 1 <= len(last_name) <= 100:
                errors.append(
                    f"Row {row_number}: first and last names must contain 1 to 100 characters."
                )
                continue
            if isinstance(course_value, bool):
                course_id = 0
            elif isinstance(course_value, int):
                course_id = course_value
            elif isinstance(course_value, float) and course_value.is_integer():
                course_id = int(course_value)
            else:
                course_id = 0
            if course_id < 1:
                errors.append(f"Row {row_number}: course_id must be a positive integer.")
                continue
            if email in seen_emails:
                errors.append(f"Row {row_number}: duplicate email in workbook.")
                continue
            seen_emails.add(email)
            imported.append(
                StudentImportRow(
                    email=email,
                    first_name=first_name,
                    last_name=last_name,
                    course_id=course_id,
                    row_number=row_number,
                )
            )

        if errors:
            raise SpreadsheetValidationError(errors)
        if not imported:
            raise SpreadsheetValidationError(
                ["Add at least one student row below the header."]
            )
        return imported
    finally:
        workbook.close()


def build_student_import_template() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    sheet.append(list(STUDENT_IMPORT_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183153")
    for column, width in {"A": 32, "B": 20, "C": 20, "D": 14}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
