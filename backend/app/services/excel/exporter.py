from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.helper import safe_spreadsheet_value


XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _add_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[dict[str, Any]],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183153")
    for record in rows:
        sheet.append(
            [safe_spreadsheet_value(record.get(header)) for header in headers]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        values = [str(header)] + [
            str(record.get(header) or "") for record in rows[:100]
        ]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(max(map(len, values)) + 2, 12), 48
        )


def export_report_workbook(report_type: str, report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    if report_type == "activities":
        summary = [
            {"metric": key.replace("_", " ").title(), "value": value}
            for key, value in report.items()
            if not isinstance(value, list)
        ]
        _add_sheet(workbook, "Summary", ["metric", "value"], summary)
        _add_sheet(
            workbook,
            "Categories",
            ["category", "count"],
            report["by_category"],
        )
        _add_sheet(
            workbook,
            "Statuses",
            ["status", "count"],
            report["by_status"],
        )
    elif report_type == "academics":
        summary = [
            {"metric": key.replace("_", " ").title(), "value": value}
            for key, value in report.items()
            if not isinstance(value, list)
        ]
        _add_sheet(workbook, "Summary", ["metric", "value"], summary)
        _add_sheet(
            workbook,
            "Attendance",
            [
                "course",
                "subject_code",
                "subject",
                "sessions_recorded",
                "present",
                "absent",
                "excused",
                "attendance_percentage",
            ],
            report["attendance"],
        )
        _add_sheet(
            workbook,
            "Assessments",
            [
                "course",
                "subject_code",
                "subject",
                "assessments_recorded",
                "average_percentage",
            ],
            report["assessments"],
        )
    else:
        raise ValueError(f"Unsupported report type: {report_type}")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
