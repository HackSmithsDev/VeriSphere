import asyncio
from io import BytesIO
from urllib.parse import parse_qs, urlparse

import httpx
import pytest
from openpyxl import load_workbook

from app.core import config
from app.services.email import sender
from app.services.email.templates import review_notification_html
from app.services.portfolio.resume_builder import build_resume_html
from app.services.telemetry import github
from app.services.telemetry import linkedin
from app.services.excel.exporter import export_report_workbook
from app.services.excel.parser import (
    SpreadsheetValidationError,
    build_student_import_template,
    parse_student_workbook,
)
from app.utils import redis as redis_cache


def test_activity_workbook_escapes_formula_cells() -> None:
    content = export_report_workbook(
        "activities",
        {
            "total_activities": 1,
            "verified_activities": 0,
            "pending_verification": 1,
            "needs_revision": 0,
            "by_category": [{"category": "=HYPERLINK(\"bad\")", "count": 1}],
            "by_status": [{"status": "pending_verification", "count": 1}],
        },
    )

    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook["Categories"]["A2"].value == "'=HYPERLINK(\"bad\")"


def test_student_import_template_and_parser() -> None:
    content = build_student_import_template()
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    assert tuple(cell.value for cell in sheet[1]) == (
        "email",
        "first_name",
        "last_name",
        "course_id",
    )
    sheet.append(["STUDENT@example.com", "Sam", "Student", 4])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    imported = parse_student_workbook(output.getvalue())

    assert len(imported) == 1
    assert imported[0].email == "student@example.com"
    assert imported[0].course_id == 4
    assert imported[0].row_number == 2


def test_student_import_rejects_formula_cells() -> None:
    workbook = load_workbook(BytesIO(build_student_import_template()))
    workbook.active.append(["student@example.com", "=1+1", "Student", 4])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(SpreadsheetValidationError, match="invalid rows") as error:
        parse_student_workbook(output.getvalue())

    assert "formulas are not allowed" in error.value.errors[0]


def test_linkedin_authorization_url_contains_signed_state(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(linkedin, "LINKEDIN_CLIENT_ID", "client-id")
    monkeypatch.setattr(linkedin, "LINKEDIN_CLIENT_SECRET", "client-secret")
    monkeypatch.setattr(
        linkedin, "LINKEDIN_REDIRECT_URI", "https://app.example.com/callback"
    )
    monkeypatch.setattr(linkedin, "LINKEDIN_SCOPES", "openid profile")

    url = linkedin.create_authorization_url(42)
    query = parse_qs(urlparse(url).query)

    assert url.startswith(linkedin.AUTHORIZE_URL)
    assert query["client_id"] == ["client-id"]
    assert query["redirect_uri"] == ["https://app.example.com/callback"]
    assert query["scope"] == ["openid profile"]
    assert linkedin.read_authorization_state(query["state"][0]) == 42


def test_linkedin_code_exchange_uses_configured_credentials_and_allowlists_profile(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(linkedin, "LINKEDIN_CLIENT_ID", "client-id")
    monkeypatch.setattr(linkedin, "LINKEDIN_CLIENT_SECRET", "client-secret")
    monkeypatch.setattr(
        linkedin, "LINKEDIN_REDIRECT_URI", "https://app.example.com/callback"
    )
    requests: list[httpx.Request] = []

    def respond(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        if request.url.path.endswith("accessToken"):
            return httpx.Response(200, json={"access_token": "provider-token"})
        return httpx.Response(
            200,
            json={
                "sub": "provider-id",
                "name": "LinkedIn Student",
                "picture": "https://media.example.com/photo",
                "email": "private@example.com",
            },
        )

    original_client = httpx.AsyncClient
    transport = httpx.MockTransport(respond)
    monkeypatch.setattr(
        linkedin.httpx,
        "AsyncClient",
        lambda **kwargs: original_client(transport=transport, **kwargs),
    )

    profile = asyncio.run(linkedin.exchange_code_for_profile("authorization-code"))

    assert profile["name"] == "LinkedIn Student"
    assert profile["picture"] == "https://media.example.com/photo"
    assert "email" not in profile
    assert requests[0].content.decode().find("client_secret=client-secret") >= 0
    assert requests[1].headers["Authorization"] == "Bearer provider-token"


def test_resume_contains_only_given_verified_items_and_escapes_html() -> None:
    html = build_resume_html(
        {
            "student": {
                "name": "<script>alert(1)</script>",
                "course": "Computer Science",
                "semester": 2,
            },
            "academic_milestones": [
                {
                    "title": "Verified <project>",
                    "subject": "Programming",
                    "semester": 1,
                }
            ],
            "activities": [],
        }
    )

    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "Verified &lt;project&gt;" in html
    assert "No verified activities yet." in html


def test_review_email_template_escapes_untrusted_content() -> None:
    html = review_notification_html(
        "<Student>",
        "<Proof>",
        "needs_revision",
        "<script>unsafe</script>",
    )

    assert "<script>" not in html
    assert "&lt;script&gt;unsafe&lt;/script&gt;" in html
    assert "&lt;Student&gt;" in html


def test_review_email_uses_configured_smtp_message(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(config, "MAIL_FROM", "noreply@example.com")
    monkeypatch.setattr(config, "MAIL_SERVER", "smtp.example.com")
    monkeypatch.setattr(config, "MAIL_USE_SSL", True)
    monkeypatch.setattr(config, "MAIL_USE_TLS", False)
    messages = []
    monkeypatch.setattr(sender, "_send_message", messages.append)

    asyncio.run(
        sender.send_review_notification(
            "student@example.com",
            "Student",
            "Milestone\nInjection",
            "verified",
            "Good work.",
        )
    )

    assert messages[0]["To"] == "student@example.com"
    assert "\n" not in messages[0]["Subject"]
    assert "Milestone Injection" in messages[0]["Subject"]
    parts = list(messages[0].iter_parts())
    assert len(parts) == 2
    assert "Milestone\nInjection" in parts[0].get_content()
    assert "Milestone\nInjection" in parts[1].get_content()


def test_github_profile_fetch_is_cached_and_fields_are_allowlisted(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    requests: list[httpx.Request] = []

    def respond(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        return httpx.Response(
            200,
            json={
                "login": "example-user",
                "name": "Example User",
                "bio": "Student",
                "public_repos": 4,
                "followers": 2,
                "following": 3,
                "html_url": "https://github.com/example-user",
                "email": "private@example.com",
            },
        )

    original_client = httpx.AsyncClient
    transport = httpx.MockTransport(respond)
    monkeypatch.setattr(
        github.httpx,
        "AsyncClient",
        lambda **kwargs: original_client(transport=transport, **kwargs),
    )
    cache_calls: list[tuple[str, int]] = []

    async def load_through_cache(key: str, ttl_seconds: int, loader):
        cache_calls.append((key, ttl_seconds))
        return await loader()

    monkeypatch.setattr(github, "get_or_load_json", load_through_cache)

    result = asyncio.run(github.get_github_profile("example-user"))

    assert result["public_repos"] == 4
    assert "email" not in result
    assert requests[0].url.path == "/users/example-user"
    assert cache_calls == [("verisphere:github:profile:example-user", 900)]


def test_github_profile_rejects_invalid_usernames_before_network() -> None:
    with pytest.raises(ValueError, match="valid public GitHub username"):
        asyncio.run(github.get_github_profile("../private"))


def test_redis_connection_errors_fall_back_to_loader(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(redis_cache, "REDIS_URL", "redis://127.0.0.1:1/0")
    calls = []

    async def loader():
        calls.append("loaded")
        return {"value": 1}

    result = asyncio.run(redis_cache.get_or_load_json("test:key", 1, loader))

    assert result == {"value": 1}
    assert calls == ["loaded"]
